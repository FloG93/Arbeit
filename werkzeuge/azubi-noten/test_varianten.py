# -*- coding: utf-8 -*-
"""Prüft die Zusatzfunktionen, indem Varianten der Mappe durchgerechnet werden.

Jede Variante ändert genau eine Stellschraube, lässt LibreOffice neu rechnen
und vergleicht das Ergebnis mit dem von Hand nachgerechneten Sollwert.

Aufruf:  python3 test_varianten.py <gebaute-datei.xlsx> <arbeitsordner>
"""

import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

def finde_recalc():
    """recalc.py rechnet die Mappe mit LibreOffice durch.

    Der Pfad steckt nicht fest im Skript: er kommt aus XLSX_RECALC oder wird
    unter ~/.claude/skills gesucht, damit das Skript auch in einer anderen
    Sitzung läuft.
    """
    aus_umgebung = os.environ.get("XLSX_RECALC")
    if aus_umgebung and Path(aus_umgebung).is_file():
        return aus_umgebung
    for wurzel in (Path.home() / ".claude" / "skills", Path("/root/.claude")):
        if wurzel.is_dir():
            treffer = sorted(wurzel.glob("**/xlsx/scripts/recalc.py"))
            if treffer:
                return str(treffer[0])
    sys.exit("recalc.py nicht gefunden. Bitte XLSX_RECALC auf den Pfad "
             "setzen (Teil der xlsx-Skill, braucht LibreOffice).")


RECALC = None

sys.path.insert(0, str(Path(__file__).resolve().parent))
import blaetter_basis as bb                                   # noqa: E402

HJ_G = 590           # Band „Gesamtprozent je Halbjahr“ im Blatt Daten
BERGER, NOVAK = bb.az_sp(0), bb.az_sp(3)      # Spalten der beiden Azubis


def krit(k):
    """Zeile des k-ten Bewertungskriteriums (0-basiert)."""
    return bb.U_P0 + k


def ohne_dollar(adr):
    return adr.replace("$", "")


def rechne(pfad):
    out = subprocess.run([sys.executable, RECALC, str(pfad), "420"],
                         capture_output=True, text=True)
    daten = json.loads(out.stdout)
    return daten


def hole(pfad, zellen):
    wb = load_workbook(pfad, data_only=True)
    werte = {}
    for name, (blatt, ref) in zellen.items():
        werte[name] = wb[blatt][ref].value
    wb.close()
    return werte


def variante(quelle, ziel, aenderungen):
    shutil.copy(quelle, ziel)
    wb = load_workbook(ziel)
    for blatt, ref, wert in aenderungen:
        wb[blatt][ref] = wert
    wb.save(ziel)
    wb.close()
    return rechne(ziel)


def nur_diagramm(erg):
    """True, wenn höchstens die geplanten Lückenmarker gemeldet wurden."""
    zusammen = erg.get("error_summary", {})
    if set(zusammen) - {"#N/A"}:
        return False
    for ort in zusammen.get("#N/A", {}).get("locations", []):
        blatt, _, ref = ort.partition("!")
        if blatt != "Azubi-Ansicht":
            return False
        zeile = int("".join(z for z in ref if z.isdigit()) or 0)
        spalte = "".join(z for z in ref if z.isalpha())
        if zeile < 81 or spalte not in ("C", "D"):
            return False
    return True


def nah(a, b, eps=0.0005):
    return a is not None and abs(a - b) < eps


def main():
    global RECALC
    RECALC = finde_recalc()
    quelle = Path(sys.argv[1])
    arbeit = Path(sys.argv[2])
    arbeit.mkdir(parents=True, exist_ok=True)
    fehler = []

    # Berger bekommt im Nachholtermin (Übung 04) einen schlechteren Wert,
    # damit sich letzte / beste / keine unterscheiden lassen.
    berger_wdh = [("Übung 04", f"{BERGER}{krit(k)}", 6) for k in range(5)]
    # 30 von 50 Punkten = 60 %, das Original hatte 88 %.

    faelle = [
        ("Wiederholung: letzte zählt (Standard)",
         berger_wdh,
         {"hj4": ("Daten", f"F{HJ_G}")},                 # HJ4 Gesamt, Azubi 1
         lambda w: nah(w["hj4"], (0.60 + 0.78) / 2)),
        ("Wiederholung: beste zählt",
         berger_wdh + [("Übung 04", f"{BERGER}{bb.U_WERT}", "beste")],
         {"hj4": ("Daten", f"F{HJ_G}")},
         lambda w: nah(w["hj4"], (0.88 + 0.78) / 2)),
        ("Wiederholung: keine Wiederholungswertung",
         berger_wdh + [("Übung 04", f"{BERGER}{bb.U_WERT}", "keine")],
         {"hj4": ("Daten", f"F{HJ_G}")},
         lambda w: nah(w["hj4"], (0.88 + 0.78) / 2)),
        ("Gewicht 3 für den Theorietest",
         [("Übung 02", ohne_dollar(bb.Z_GEWICHT), 3)],
         {"hj4": ("Daten", f"F{HJ_G}")},
         lambda w: nah(w["hj4"], (0.88 * 1 + 0.78 * 3) / 4)),
        ("Gewichtung Praxis 70 / Theorie 30 eingeschaltet",
         [("Einstellungen", "C18", "ja")],
         {"hj4": ("Daten", f"F{HJ_G}")},
         lambda w: nah(w["hj4"], 0.88 * 0.7 + 0.78 * 0.3)),
        ("Notengrenze für die 2 auf 85 % angehoben",
         [("Einstellungen", "C8", 0.85)],
         {"note": ("Gesprächsseite", "E17")},            # Gesamt-Note
         lambda w: w["note"] == "3 – befriedigend"),
        ("Schwelle „wenig Datenbasis“ auf 2 gesenkt",
         [("Einstellungen", "C17", 2)],
         {"hinweis": ("Gesprächsseite", "F17")},
         lambda w: (w["hinweis"] or "") == ""),
        ("Filter auf Industriemechaniker/-in",
         [("Azubi-Ansicht", "C4", "Industriemechaniker/-in")],
         {"anzahl": ("Daten", "B623"), "erster": ("Daten", "D660")},
         lambda w: w["anzahl"] == 2 and w["erster"] == "Hoffmann, Jonas"),
        ("Azubi auf ausgeschieden gesetzt",
         [("Stammdaten", "G6", "ausgeschieden")],
         {"anzahl": ("Daten", "B623"),
          "uebung1": ("Übungen", "L6")},                 # bewertete Azubis
         lambda w: w["anzahl"] == 3 and w["uebung1"] == 2),
        ("Ganze Zeile n.b. ohne Grund",
         [("Übung 03", f"{BERGER}{krit(k)}", "n.b.") for k in range(4)]
         + [("Übung 03", ohne_dollar(bb.Z_STATUS), "abgeschlossen")],
         {"pruef": ("Start", "M16"), "ampel": ("Start", "J6")},
         lambda w: w["pruef"] == 1 and str(w["ampel"]).startswith("GELB")),
        ("Unvollständig ausgefüllte Zeile",
         [("Übung 03", f"{NOVAK}{krit(0)}", 5)],
         {"pruef": ("Start", "M15"), "ampel": ("Start", "J6")},
         lambda w: w["pruef"] == 1 and str(w["ampel"]).startswith("ROT")),
        ("Gesprächsseite bleibt vom Filter der Azubi-Ansicht unabhängig",
         [("Azubi-Ansicht", "C4", "Industriemechaniker/-in")],
         {"gefiltert": ("Daten", "B623"), "alle": ("Daten", "B819"),
          "erster_alle": ("Daten", "D820")},
         lambda w: w["gefiltert"] == 2 and w["alle"] == 4
         and w["erster_alle"] == "Berger, Lukas"),
        ("Zwei Azubis mit gleichem Namen",
         [("Stammdaten", "B9", "Berger"), ("Stammdaten", "C9", "Lukas")],
         {"pruef": ("Start", "M18"), "ampel": ("Start", "J6")},
         lambda w: w["pruef"] == 2 and str(w["ampel"]).startswith("ROT")),
        ("Übung ohne Datum",
         [("Übung 03", ohne_dollar(bb.Z_DATUM), None)],
         {"pruef": ("Start", "M12"), "ampel": ("Start", "J6")},
         lambda w: w["pruef"] == 1 and str(w["ampel"]).startswith("ROT")),
    ]

    for i, (name, aenderungen, zellen, pruefung) in enumerate(faelle, 1):
        ziel = arbeit / f"v{i:02d}.xlsx"
        erg = variante(quelle, ziel, aenderungen)
        # Erlaubt sind nur die gewollten #N/A in den Diagramm-Hilfszellen:
        # sie sorgen dafür, dass die Linie dort eine Lücke lässt.
        fehlerfrei = nur_diagramm(erg)
        werte = hole(ziel, zellen)
        ok = pruefung(werte)
        status = "OK  " if (ok and fehlerfrei) else "FEHL"
        print(f"{status} {i:2}. {name}")
        print(f"       {werte}  (nur geplante #N/A: {fehlerfrei}, gesamt {erg.get('total_errors')})")
        if not (ok and fehlerfrei):
            fehler.append(name)

    print()
    if fehler:
        print(f"{len(fehler)} Fall/Fälle nicht bestanden:")
        for n in fehler:
            print("  -", n)
        sys.exit(1)
    print(f"Alle {len(faelle)} Fälle bestanden.")


if __name__ == "__main__":
    main()
