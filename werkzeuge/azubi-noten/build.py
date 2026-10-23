# -*- coding: utf-8 -*-
"""Baut die Datei „Azubi-Leistungen.xlsx“.

Aufruf:  python3 build.py [zieldatei]

Die Mappe ist reines .xlsx: keine Makros, keine Verknüpfungen nach außen,
keine Funktionen jünger als Excel 2007 – damit sie auf SharePoint im Browser,
in älteren Excel-Versionen und in LibreOffice gleich funktioniert.
"""

import datetime
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment

import blaetter_basis as bb
from stil import (BG, EINGABE_BG, FMT_DATUM, LINKS, MAX_AZUBIS, N_UEBUNGEN,
                  NAVY_500, NAVY_700, TEAL, TXT_HELL, UEB_BLATT, blattkopf,
                  breiten, f, setze, verbinde)

ZIEL = "Azubi-Leistungen.xlsx"

# Blätter in der Reihenfolge, in der sie in der Mappe stehen sollen
NAV_BLAETTER = ["Start", "Azubi-Ansicht", "Gesprächsseite", "Halbjahr",
                "Übungen", "Stammdaten", "Einstellungen", "Anleitung"]


def neue_mappe():
    wb = Workbook()
    wb.remove(wb.active)
    for name in NAV_BLAETTER:
        wb.create_sheet(name)
    for i in range(1, N_UEBUNGEN + 1):
        wb.create_sheet(UEB_BLATT.format(i))
    wb.create_sheet("Vorlage")
    wb.create_sheet("Daten")
    wb.properties.title = "Azubi-Leistungen"
    wb.properties.creator = "Ausbildung"
    wb.properties.description = ("Erfassung und Auswertung der Leistungen "
                                 "von Auszubildenden")
    wb.calculation.fullCalcOnLoad = True
    return wb


def platzhalter(ws, titel, untertitel, letzte="N"):
    """Vorläufiges Blatt, damit die Verlinkung schon in Stufe 1 trägt."""
    blattkopf(ws, letzte, titel, untertitel,
              zurueck=(ws.title != "Start"))
    verbinde(ws, "A4:H5", "Dieses Blatt entsteht in der nächsten Stufe.",
             font=f(11, color=TXT_HELL), bg=BG, align=LINKS)


# ------------------------------------------------------------ Beispieldaten --
AZUBIS = [
    ("Berger", "Lukas", "Elektroniker/-in für Betriebstechnik",
     datetime.date(2024, 9, 1), "aktiv"),
    ("Demir", "Elif", "Elektroniker/-in für Betriebstechnik",
     datetime.date(2024, 9, 1), "aktiv"),
    ("Hoffmann", "Jonas", "Industriemechaniker/-in",
     datetime.date(2025, 9, 1), "aktiv"),
    ("Novak", "Mira", "Industriemechaniker/-in",
     datetime.date(2025, 9, 1), "aktiv"),
]

BEISPIEL_UEBUNGEN = [
    dict(
        nr=1, name="Bohren und Senken", art="Praxis",
        datum=datetime.date(2026, 3, 12), status="abgeschlossen", wer="MS",
        params=[("Maßhaltigkeit", "Alle Maße in Toleranz ±0,1 mm", 10),
                ("Bohrungen korrekt", "10 von 10 Bohrungen korrekt", 10),
                ("Senkungen", "Senkwinkel und Tiefe gleichmäßig", 10),
                ("Oberfläche", "Grat entfernt, keine Riefen", 10),
                ("Arbeitssicherheit", "PSA getragen, Maschine gesichert", 10)],
        werte=[[9, 10, 8, 7, 10],
               [10, 10, 9, 9, 10],
               ["n.b.", "n.b.", "n.b.", "n.b.", "n.b."],
               [6, 5, 7, 6, 8]],
        gruende=["", "", "krank", ""],
        bemerkungen=["Sauber gearbeitet, etwas langsam.", "Vorbildlich.",
                     "Nachholtermin vereinbart.",
                     "Bohrbild üben, Sicherheit war in Ordnung."],
    ),
    dict(
        nr=2, name="Grundlagen Elektrotechnik", art="Theorie",
        datum=datetime.date(2026, 4, 24), status="abgeschlossen", wer="AB",
        params=[("Ohmsches Gesetz", "6 von 6 Aufgaben richtig", 10),
                ("Reihenschaltung", "5 von 5 Aufgaben richtig", 10),
                ("Parallelschaltung", "5 von 5 Aufgaben richtig", 10),
                ("Leistungsberechnung", "10 von 10 Aufgaben richtig", 20)],
        werte=[[8, 9, 7, 15],
               [9, 10, 9, 18],
               [7, 6, 8, 12],
               [5, 6, "n.b.", 10]],
        gruende=["", "", "", "Berufsschule"],
        bemerkungen=["", "", "Leistungsberechnung wiederholen.",
                     "Aufgabenteil fiel wegen Berufsschule aus."],
    ),
    dict(
        nr=3, name="Leitungen verlegen", art="Praxis",
        datum=datetime.date(2026, 10, 2), status="laufend", wer="MS",
        params=[("Verlegeart", "Vorgabe eingehalten", 10),
                ("Befestigung", "Abstände nach Vorgabe", 10),
                ("Biegeradien", "Mindestradien eingehalten", 10),
                ("Ordnung/Sauberkeit", "Arbeitsplatz aufgeräumt", 10)],
        werte=[[10, 9, 9, 10],
               [8, 8, 9, 8],
               [6, 5, 6, 6],
               [None, None, None, None]],
        gruende=["", "", "", ""],
        bemerkungen=["", "", "Nach der Krankheit noch unsicher.",
                     "Noch nicht bewertet."],
    ),
    dict(
        nr=4, name="Bohren und Senken – Nachholtermin", art="Praxis",
        datum=datetime.date(2026, 6, 9), status="abgeschlossen", wer="MS",
        wdh_von="Bohren und Senken",
        params=[("Maßhaltigkeit", "Alle Maße in Toleranz ±0,1 mm", 10),
                ("Bohrungen korrekt", "10 von 10 Bohrungen korrekt", 10),
                ("Senkungen", "Senkwinkel und Tiefe gleichmäßig", 10),
                ("Oberfläche", "Grat entfernt, keine Riefen", 10),
                ("Arbeitssicherheit", "PSA getragen, Maschine gesichert", 10)],
        werte=[[None] * 5,
               [None] * 5,
               [8, 9, 8, 7, 9],
               [None] * 5],
        gruende=["", "", "", ""],
        bemerkungen=["", "", "Nachholtermin nach Krankheit.", ""],
    ),
]


def beispieldaten(wb):
    sd = wb["Stammdaten"]
    for i, (name, vorname, beruf, start, status) in enumerate(AZUBIS):
        r = bb.SD0 + i
        sd[f"B{r}"] = name
        sd[f"C{r}"] = vorname
        sd[f"D{r}"] = beruf
        sd[f"F{r}"] = start
        sd[f"F{r}"].number_format = FMT_DATUM
        sd[f"G{r}"] = status

    uv = wb["Übungen"]
    for ue in BEISPIEL_UEBUNGEN:
        uv[f"C{bb.UV0 + ue['nr'] - 1}"] = ue["name"]
        ws = wb[UEB_BLATT.format(ue["nr"])]
        ws[f"J{bb.R_K1}"] = ue["art"]
        ws[f"N{bb.R_K1}"] = ue["datum"]
        ws[f"N{bb.R_K1}"].number_format = FMT_DATUM
        ws[f"C{bb.R_K2}"] = ue["status"]
        ws[f"J{bb.R_K2}"] = ue["wer"]
        if ue.get("wdh_von"):
            ws[f"C{bb.R_K3}"] = ue["wdh_von"]
        for j, (pname, phor, pmax) in enumerate(ue["params"]):
            sp = bb.PARAM_SPALTEN[j]
            ws[f"{sp}{bb.R_PNAME}"] = pname
            ws[f"{sp}{bb.R_PHOR}"] = phor
            if pmax != 10:
                # abweichende Höchstpunktzahl überschreibt den Vorschlag
                ws[f"{sp}{bb.R_PMAX}"] = pmax
        for i, zeile in enumerate(ue["werte"]):
            r = bb.R_DAT0 + i
            for j, wert in enumerate(zeile):
                if wert is not None and wert != "":
                    ws[f"{bb.PARAM_SPALTEN[j]}{r}"] = wert
            if ue["gruende"][i]:
                ws[f"P{r}"] = ue["gruende"][i]
            if ue["bemerkungen"][i]:
                ws[f"Q{r}"] = ue["bemerkungen"][i]

    # Startauswahl, damit Azubi-Ansicht und Diagramm nicht leer wirken
    wb["Azubi-Ansicht"]["C5"] = f"{AZUBIS[0][0]}, {AZUBIS[0][1]}"
    wb["Gesprächsseite"]["C6"] = f"{AZUBIS[0][0]}, {AZUBIS[0][1]}"


# ------------------------------------------------------------------- Bauen ---
def main():
    ziel = sys.argv[1] if len(sys.argv) > 1 else ZIEL
    wb = neue_mappe()

    bb.baue_einstellungen(wb["Einstellungen"])
    bb.einstellungen_namen(wb)
    bb.baue_stammdaten(wb["Stammdaten"])
    bb.baue_uebungsliste(wb["Übungen"])
    for i in range(1, N_UEBUNGEN + 1):
        bb.baue_uebungsblatt(wb[UEB_BLATT.format(i)], i)
    bb.baue_uebungsblatt(wb["Vorlage"], None)

    import blaetter_auswertung as ba
    ba.baue_daten(wb["Daten"])
    ba.daten_namen(wb)
    ba.patch_uebungsliste(wb["Übungen"])
    ba.baue_start(wb["Start"])
    ba.baue_azubi(wb["Azubi-Ansicht"])
    ba.baue_halbjahr(wb["Halbjahr"])

    import blaetter_extra as be
    be.baue_gespraech(wb["Gesprächsseite"])
    be.baue_anleitung(wb["Anleitung"])

    beispieldaten(wb)
    # Das Rechenblatt wird ausgeblendet: es enthält nur Zwischenergebnisse
    # und ist so vor versehentlichem Überschreiben geschützt.
    wb["Daten"].sheet_state = "hidden"
    wb.active = 0
    wb.save(ziel)
    print(f"geschrieben: {ziel}")


if __name__ == "__main__":
    main()
