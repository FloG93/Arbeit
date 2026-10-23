# -*- coding: utf-8 -*-
"""Gemeinsame Gestaltung und kleine Helfer für die Azubi-Leistungsmappe.

Alle Farben, Schriften und Zellstile stehen hier an einer Stelle, damit die
Blätter untereinander gleich aussehen. Wer das Aussehen ändern will, ändert es
hier — nicht verstreut in den einzelnen Blattbauten.
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

# ---------------------------------------------------------------- Grundmaße --
MAX_AZUBIS = 25          # Zeilen je Übungsblatt
MAX_PARAM = 8            # Bewertungsparameter je Übung
N_UEBUNGEN = 20          # vorbereitete Übungsblätter
MAX_HJ = 8               # dargestellte Halbjahre (3,5 Jahre Ausbildung = 7)
N_WARN = 10              # Zeilen der Warnübersicht auf der Startseite

SCHRIFT = "Arial"

# ------------------------------------------------------------------ Farben ---
# Eine ruhige, dunkelblaue Grundfarbe trägt die Kopfbereiche; Türkis setzt
# Akzente. Die Notenfarben folgen der Vorgabe dunkelgrün → dunkelrot.
NAVY_900 = "16263B"
NAVY_700 = "22395C"
NAVY_500 = "34557F"
NAVY_100 = "DCE4EF"
TEAL = "0E7C86"
TEAL_100 = "D8EDEF"
BG = "F4F6F9"
PANEL = "FFFFFF"
LINIE = "D4DAE3"
LINIE_STARK = "9AA7B8"
TXT = "1C2733"
TXT_HELL = "6B7684"
TXT_WEISS = "FFFFFF"

EINGABE_BG = "FFF8E1"        # hellgelb: hier wird getippt
EINGABE_RAND = "E0C068"
FORMEL_BG = "EDF1F6"         # hellgrau-blau: rechnet die Datei selbst
FORMEL_TXT = "3F5168"
GESPERRT_BG = "E8EAED"       # Blatt noch nicht freigeschaltet

AMPEL_GRUEN = "1E7B34"
AMPEL_GELB = "E8A33D"
AMPEL_ROT = "C0392B"

# Note 1..6: (Hintergrund, Schriftfarbe)
NOTENFARBEN = [
    ("1E7B34", "FFFFFF"),   # 1 sehr gut        – dunkelgrün
    ("A8D08D", "1C2733"),   # 2 gut             – hellgrün
    ("FFD966", "1C2733"),   # 3 befriedigend    – gelb
    ("F4A742", "1C2733"),   # 4 ausreichend     – orange
    ("E03C31", "FFFFFF"),   # 5 mangelhaft      – rot
    ("8E1B13", "FFFFFF"),   # 6 ungenügend      – dunkelrot
]
NOTENTEXTE = ["sehr gut", "gut", "befriedigend", "ausreichend",
              "mangelhaft", "ungenügend"]
NOTENGRENZEN = [0.92, 0.81, 0.67, 0.50, 0.30, 0.00]

# --------------------------------------------------------------- Zahlformate -
FMT_PROZ = "0.0%"
FMT_PUNKT = "0.##"
FMT_DATUM = "DD.MM.YYYY"
FMT_NOTE = "0.0"
FMT_HJ = '0;0;"–"'       # 0 bedeutet „nicht bestimmbar“ und zeigt einen Strich


# ------------------------------------------------------------------ Helfer ---
def f(size=11, bold=False, color=TXT, italic=False, underline=None):
    return Font(name=SCHRIFT, size=size, bold=bold, color=color,
                italic=italic, underline=underline)


def fill(color):
    return PatternFill("solid", fgColor=color)


def kante(color=LINIE, stil="thin"):
    return Side(style=stil, color=color)


def rahmen(color=LINIE, stil="thin"):
    s = kante(color, stil)
    return Border(left=s, right=s, top=s, bottom=s)


RAHMEN_ZART = rahmen(LINIE)
RAHMEN_EINGABE = rahmen(EINGABE_RAND)
MITTE = Alignment(horizontal="center", vertical="center", wrap_text=False)
MITTE_UM = Alignment(horizontal="center", vertical="center", wrap_text=True)
LINKS = Alignment(horizontal="left", vertical="center")
LINKS_UM = Alignment(horizontal="left", vertical="center", wrap_text=True)
LINKS_OBEN = Alignment(horizontal="left", vertical="top", wrap_text=True)
RECHTS = Alignment(horizontal="right", vertical="center")


def setze(ws, ref, wert=None, font=None, bg=None, align=None, border=None,
          fmt=None):
    """Schreibt Wert und Stil in eine Zelle (bei verbundenen nur links oben)."""
    c = ws[ref]
    if wert is not None:
        c.value = wert
    if font is not None:
        c.font = font
    if bg is not None:
        c.fill = fill(bg)
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    if fmt is not None:
        c.number_format = fmt
    return c


def bereich(ws, rng, font=None, bg=None, align=None, border=None, fmt=None):
    """Stil über einen Zellbereich legen, ohne Werte anzufassen."""
    for row in ws[rng]:
        for c in row:
            if font is not None:
                c.font = font
            if bg is not None:
                c.fill = fill(bg)
            if align is not None:
                c.alignment = align
            if border is not None:
                c.border = border
            if fmt is not None:
                c.number_format = fmt


def verbinde(ws, rng, wert=None, **kw):
    ws.merge_cells(rng)
    links_oben = rng.split(":")[0]
    return setze(ws, links_oben, wert, **kw)


def aussenrahmen(ws, rng, color=LINIE_STARK, stil="medium"):
    """Legt nur außen herum eine Linie — innen bleibt, was schon da ist."""
    zellen = ws[rng]
    s = kante(color, stil)
    n_r = len(zellen)
    n_c = len(zellen[0])
    for i, row in enumerate(zellen):
        for j, c in enumerate(row):
            b = c.border
            c.border = Border(
                left=s if j == 0 else b.left,
                right=s if j == n_c - 1 else b.right,
                top=s if i == 0 else b.top,
                bottom=s if i == n_r - 1 else b.bottom,
            )


def breiten(ws, mapping):
    for sp, w in mapping.items():
        ws.column_dimensions[sp].width = w


def hoehen(ws, mapping):
    for r, h in mapping.items():
        ws.row_dimensions[r].height = h


def link_zelle(ws, ref, ziel, text, font=None, bg=None, align=None,
               border=None):
    """Echter interner Hyperlink — funktioniert auch in Excel im Browser."""
    c = ws[ref]
    c.value = text
    c.hyperlink = Hyperlink(ref=ref, location=ziel)
    c.font = font or f(11, bold=True, color=TEAL, underline="single")
    if bg is not None:
        c.fill = fill(bg)
    if align is not None:
        c.alignment = align
    if border is not None:
        c.border = border
    return c


def blattkopf(ws, letzte_spalte, titel, untertitel, zurueck=True,
              tabfarbe=NAVY_700, link_breite=4):
    """Einheitliches Kopfband: Titel links, Rücksprung zur Startseite rechts."""
    ws.sheet_properties.tabColor = tabfarbe
    ws.sheet_view.showGridLines = False
    n = ws[letzte_spalte + "1"].column
    trenn = get_column_letter(max(2, n - link_breite + 1))
    vor = get_column_letter(max(1, n - link_breite))

    verbinde(ws, f"A1:{vor}1", titel,
             font=f(20, bold=True, color=TXT_WEISS), bg=NAVY_700,
             align=Alignment(horizontal="left", vertical="center", indent=1))
    verbinde(ws, f"A2:{vor}2", untertitel,
             font=f(10, color=NAVY_100), bg=NAVY_700,
             align=Alignment(horizontal="left", vertical="center", indent=1))
    bereich(ws, f"A1:{letzte_spalte}2", bg=NAVY_700)
    if zurueck:
        verbinde(ws, f"{trenn}1:{letzte_spalte}2")
        link_zelle(ws, f"{trenn}1", "Start!A1", "«  Zurück zur Startseite",
                   font=f(11, bold=True, color=TXT_WEISS, underline="single"),
                   bg=NAVY_700, align=MITTE)
    hoehen(ws, {1: 34, 2: 18})


def legende_eingabe(ws, rng):
    """Kurze Erklärung der Zellfarben — steht auf jedem Blatt."""
    verbinde(
        ws, rng,
        "Farben:  gelb hinterlegt = hier eintragen   ·   grau-blau = rechnet "
        "die Datei selbst (bitte nicht überschreiben)   ·   „n.b.“ = nicht "
        "bewertet",
        font=f(9.5, color=TXT_HELL), bg=BG, align=LINKS_UM)


def notenlegende(ws, start_zelle, waagrecht=True, mit_grenzen=True):
    """Farblegende Note 1–6. Zeigt immer Zahl + Kurztext, nie nur Farbe."""
    sp = ws[start_zelle].column
    zl = ws[start_zelle].row
    for i in range(6):
        bgc, txt = NOTENFARBEN[i]
        r = zl if waagrecht else zl + i
        c = sp + i if waagrecht else sp
        ref = f"{get_column_letter(c)}{r}"
        if mit_grenzen:
            # Grenze kommt aus den Einstellungen, damit die Legende mitwandert
            wert = (f'=INDEX(Noten,{i + 1})&" – "&INDEX(NText,{i + 1})&'
                    f'" (ab "&TEXT(INDEX(NGrenze,{i + 1})*100,"0")&" %)"')
        else:
            wert = f'=INDEX(Noten,{i + 1})&" – "&INDEX(NText,{i + 1})'
        setze(ws, ref, wert, font=f(9.5, bold=True, color=txt), bg=bgc,
              align=MITTE, border=rahmen(PANEL))
    return zl


def notenlegende_bereiche(ws, bereiche, mit_grenzen=True):
    """Farblegende in sechs vorgegebene Bereiche – für schmale Spalten."""
    for i, rng in enumerate(bereiche):
        bgc, txt = NOTENFARBEN[i]
        if mit_grenzen:
            wert = (f'=INDEX(Noten,{i + 1})&" – "&INDEX(NText,{i + 1})&'
                    f'" (ab "&TEXT(INDEX(NGrenze,{i + 1})*100,"0")&" %)"')
        else:
            wert = f'=INDEX(Noten,{i + 1})&" – "&INDEX(NText,{i + 1})'
        verbinde(ws, rng, wert, font=f(9.5, bold=True, color=txt), bg=bgc,
                 align=MITTE, border=rahmen(PANEL))


def note_index(pct_ref):
    """Zeilennummer im Notenschlüssel für einen Prozentwert (0–1)."""
    return (f'MAX(1,MIN(COUNT(NGrenze),COUNT(NGrenze)+1-'
            f'COUNTIF(NGrenze,"<="&{pct_ref})))')


def note_zahl(pct_ref):
    return f"INDEX(Noten,{note_index(pct_ref)})"


def note_text(pct_ref):
    idx = note_index(pct_ref)
    return f'INDEX(Noten,{idx})&" – "&INDEX(NText,{idx})'


def datum_text(ref):
    """Datum als Text, ohne TEXT() — dessen Formatcodes sind sprachabhängig."""
    return (f'RIGHT("0"&DAY({ref}),2)&"."&RIGHT("0"&MONTH({ref}),2)&"."&'
            f'YEAR({ref})')


def ampel_regeln(ws, rng, pct_ref):
    """Bedingte Formatierung Note 1–6 für einen Bereich.

    Die Grenzen werden über INDEX(NGrenze,k) aus dem Einstellungsblatt geholt —
    ändert dort jemand eine Grenze, färben sich die Zellen sofort neu.
    """
    from openpyxl.formatting.rule import FormulaRule
    for i in range(6):
        bgc, txt = NOTENFARBEN[i]
        regel = FormulaRule(
            formula=[f'AND(ISNUMBER({pct_ref}),{pct_ref}>=INDEX(NGrenze,'
                     f'{i + 1}))'],
            fill=fill(bgc), font=f(11, bold=True, color=txt),
            stopIfTrue=True)
        ws.conditional_formatting.add(rng, regel)


def druck(ws, bereich_, quer=True, hoch_seiten=0, kopfzeilen=None):
    """Druckbild festlegen: A4, auf Seitenbreite skaliert."""
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.print_area = bereich_
    ws.page_setup.orientation = "landscape" if quer else "portrait"
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = hoch_seiten
    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.4
    if kopfzeilen:
        ws.print_title_rows = kopfzeilen


UEB_BLATT = "Übung {:02d}"
