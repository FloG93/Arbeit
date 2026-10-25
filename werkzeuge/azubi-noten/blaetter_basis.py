# -*- coding: utf-8 -*-
"""Stufe 1: Einstellungen, Stammdaten, Übungsliste, Übungsblätter, Vorlage."""

from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

from stil import (BG, EINGABE_BG, EINGABE_RAND, FMT_DATUM, FMT_HJ, FMT_PROZ,
                  FMT_PUNKT, FORMEL_BG, FORMEL_TXT, GESPERRT_BG, LINIE,
                  LINIE_STARK, LINKS, LINKS_UM, MAX_AZUBIS, MAX_PARAM, MITTE,
                  MITTE_UM, N_UEBUNGEN, NAVY_100, NAVY_500, NAVY_700, PANEL,
                  RAHMEN_EINGABE, RAHMEN_ZART, TEAL, TEAL_100, TXT, TXT_HELL,
                  TXT_WEISS, UEB_BLATT, ampel_regeln, aussenrahmen, bereich,
                  blattkopf,
                  breiten, datum_text, f, fill, hoehen, legende_eingabe,
                  druck, link_zelle, note_text, notenlegende_bereiche,
                  rahmen, setze, verbinde)

# ------------------------------------------------- Raster der Übungsblätter --
# Kriterien stehen untereinander, Azubis nebeneinander. Bei vielen Kriterien
# scrollt man damit nach unten statt zur Seite. Der fixierte Kopf (bis
# einschließlich U_THEAD) hält Name, Prozent und Note sichtbar, während unten
# die Punkte eingetragen werden.
AZ0 = 5                 # Spalte E trägt den ersten Azubi
U_TITEL = 1
U_BANNER = 2
U_K1 = 3                # Übungsname | Art | Datum
U_K2 = 4                # Status | Bewertet von | Gewicht | Wiederholung von
U_NAME = 5
U_VORNAME = 6
U_PUNKTE = 7            # „44 / 50“, nur zur Anzeige
U_PROZ = 8
U_NOTE = 9
U_THEAD = 10
U_P0 = 11                                   # erste Kriterienzeile
U_P1 = U_P0 + MAX_PARAM - 1
U_SUMME = U_P1 + 2
U_MAX = U_P1 + 3
U_HJ = U_P1 + 4
U_GRUND = U_P1 + 5
U_BEM = U_P1 + 6
U_WERT = U_P1 + 7
U_STAT = U_P1 + 8       # Azubi-Status, nur damit die Einfärbung ihn findet
U_LEG = U_P1 + 10
U_FARB = U_P1 + 11
U_HINW = U_P1 + 13

# Kopffelder des Übungsblattes – jede Adresse steht genau hier
Z_NAME = "$C$3"
Z_ART = "$H$3"
Z_DATUM = "$L$3"
Z_STATUS = "$C$4"
Z_WER = "$H$4"
Z_GEWICHT = "$L$4"
Z_WDH = "$Q$4"
Z_PMAX = f"$D${U_P0}:$D${U_P1}"

SD0 = 6                                     # erste Datenzeile der Stammdaten
SD1 = SD0 + MAX_AZUBIS - 1
UV0 = 6                                     # erste Zeile der Übungsliste
UV1 = UV0 + N_UEBUNGEN - 1


def az_sp(i):
    """Spaltenbuchstabe des i-ten Azubis (0-basiert)."""
    return get_column_letter(AZ0 + i)


AZ_ERSTE = az_sp(0)
AZ_LETZTE = az_sp(MAX_AZUBIS - 1)
DRUCK_SPALTEN = min(MAX_AZUBIS, 25)         # ein Jahrgangsblock je Ausdruck


# ============================================================ Einstellungen ==
def baue_einstellungen(ws):
    blattkopf(ws, "F", "Einstellungen",
              "Zentrale Stellschrauben – wirken sofort in der ganzen Datei",
              tabfarbe=NAVY_500, link_breite=1)
    breiten(ws, {"A": 3, "B": 30, "C": 16, "D": 22, "E": 12, "F": 46})
    legende_eingabe(ws, "A3:F3")
    hoehen(ws, {3: 24, 4: 10})

    # ---- Notenschlüssel -----------------------------------------------------
    verbinde(ws, "B5:F5", "Notenschlüssel (IHK)",
             font=f(13, bold=True, color=NAVY_700), align=LINKS)
    hoehen(ws, {5: 26, 6: 22})
    kopf = ["Note", "ab Prozent", "Kurztext", "Farbe", "Bedeutung"]
    for i, t in enumerate(kopf):
        setze(ws, f"{chr(66 + i)}6", t, font=f(10, bold=True, color=TXT_WEISS),
              bg=NAVY_500, align=MITTE, border=rahmen(PANEL))
    from stil import NOTENFARBEN, NOTENGRENZEN, NOTENTEXTE
    bedeutung = ["sehr gut – Spitzenleistung", "gut – sicher über Erwartung",
                 "befriedigend – im Blick behalten",
                 "ausreichend – gefährdet, Gespräch empfohlen",
                 "mangelhaft – nicht bestanden",
                 "ungenügend – nicht bestanden"]
    for i in range(6):
        r = 7 + i
        bgc, txtc = NOTENFARBEN[i]
        setze(ws, f"B{r}", i + 1, font=f(12, bold=True), bg=EINGABE_BG,
              align=MITTE, border=RAHMEN_EINGABE)
        setze(ws, f"C{r}", NOTENGRENZEN[i], font=f(12), bg=EINGABE_BG,
              align=MITTE, border=RAHMEN_EINGABE, fmt="0%")
        setze(ws, f"D{r}", NOTENTEXTE[i], font=f(11), bg=EINGABE_BG,
              align=LINKS, border=RAHMEN_EINGABE)
        setze(ws, f"E{r}", f"Note {i + 1}", font=f(10, bold=True, color=txtc),
              bg=bgc, align=MITTE, border=rahmen(PANEL))
        setze(ws, f"F{r}", bedeutung[i], font=f(10, color=TXT_HELL),
              align=LINKS, border=RAHMEN_ZART)
        hoehen(ws, {r: 22})
    verbinde(
        ws, "B13:F13",
        "Die Grenzen gelten „ab“ (92 % → Note 1). Reihenfolge Note 1 oben bis "
        "Note 6 unten bitte beibehalten – alle Formeln und Farben lesen die "
        "Tabelle von oben nach unten.",
        font=f(9.5, color=TXT_HELL), bg=BG, align=LINKS_UM)
    hoehen(ws, {13: 28, 14: 10})

    # ---- Grundeinstellungen -------------------------------------------------
    verbinde(ws, "B15:F15", "Grundeinstellungen",
             font=f(13, bold=True, color=NAVY_700), align=LINKS)
    hoehen(ws, {15: 26})
    grund = [
        ("Punkte je Parameter (Standard)", 10, "0",
         "Wird auf neuen Übungsblättern als Höchstpunktzahl vorgeschlagen."),
        ("Schwelle „wenig Datenbasis“", 3, "0",
         "Unter so vielen bewerteten Übungen erscheint neben der Halbjahres"
         "note ein Vermerk. Die Berechnung ändert sich dadurch nicht."),
        ("Gewichtung Praxis/Theorie verwenden?", "nein", None,
         "nein = alle Übungen zählen gleich (Mittel der Prozentwerte). "
         "ja = Praxis und Theorie werden getrennt gemittelt und mit den "
         "Gewichten unten zusammengezogen."),
        ("Gewicht Praxis", 0.7, "0%",
         "Nur wirksam, wenn oben „ja“ steht."),
        ("Gewicht Theorie", 0.3, "0%",
         "Nur wirksam, wenn oben „ja“ steht."),
        ("Standard-Wertung bei Wiederholung", "letzte", None,
         "Gilt, solange auf dem Wiederholungsblatt beim Azubi nichts anderes "
         "steht: letzte / beste / keine."),
        ("Warnschwelle (Note ab …)", 4, "0",
         "Ab dieser Halbjahresnote erscheint der Azubi in der Warnübersicht "
         "auf der Startseite."),
    ]
    for i, (name, wert, fmt, hinweis) in enumerate(grund):
        r = 16 + i
        setze(ws, f"B{r}", name, font=f(11, bold=True), align=LINKS,
              border=RAHMEN_ZART)
        setze(ws, f"C{r}", wert, font=f(12, bold=True), bg=EINGABE_BG,
              align=MITTE, border=RAHMEN_EINGABE, fmt=fmt)
        verbinde(ws, f"D{r}:F{r}", hinweis, font=f(9.5, color=TXT_HELL),
                 align=LINKS_UM, border=RAHMEN_ZART)
        hoehen(ws, {r: 30})
    dv_ja = DataValidation(type="list", formula1='"ja,nein"', allow_blank=False)
    ws.add_data_validation(dv_ja)
    dv_ja.add("C18")
    dv_w = DataValidation(type="list", formula1='"letzte,beste,keine"',
                          allow_blank=False)
    ws.add_data_validation(dv_w)
    dv_w.add("C21")
    hoehen(ws, {23: 10})

    # ---- Ausbildungsberufe --------------------------------------------------
    verbinde(ws, "B24:F24", "Ausbildungsberufe (Auswahl in den Stammdaten)",
             font=f(13, bold=True, color=NAVY_700), align=LINKS)
    hoehen(ws, {24: 26})
    berufe = ["Elektroniker/-in für Betriebstechnik",
              "Industriemechaniker/-in", "Mechatroniker/-in",
              "Fachkraft für Metalltechnik"]
    for i in range(10):
        r = 25 + i
        setze(ws, f"B{r}", berufe[i] if i < len(berufe) else None,
              font=f(11), bg=EINGABE_BG, align=LINKS, border=RAHMEN_EINGABE)
        hoehen(ws, {r: 20})
    verbinde(ws, "C25:F25",
             "Freie Zeilen bleiben leer – sie erscheinen dann als leerer "
             "Eintrag in der Auswahlliste.",
             font=f(9.5, color=TXT_HELL), align=LINKS_UM)

    # ---- Hinweis Farben -----------------------------------------------------
    verbinde(ws, "B36:F38",
             "Farben ändern: Die Prozentgrenzen oben wirken sofort auf alle "
             "Einfärbungen. Die Farbtöne selbst liegen in der bedingten "
             "Formatierung der Blätter (Start > Bedingte Formatierung > "
             "Regeln verwalten). Ohne Makros lässt sich ein Farbton nicht aus "
             "einer Zelle steuern – deshalb steht neben jeder Farbe immer "
             "auch die Note als Zahl mit Kurztext.",
             font=f(10, color=TXT), bg=BG, align=LINKS_UM)
    aussenrahmen(ws, "B36:F38", LINIE_STARK, "thin")
    hoehen(ws, {36: 22, 37: 22, 38: 22})
    druck(ws, "A1:F38", quer=False)
    ws.freeze_panes = "A3"


def einstellungen_namen(wb):
    """Benannte Bereiche – machen die Formeln überall lesbar."""
    from openpyxl.workbook.defined_name import DefinedName
    namen = {
        "Noten": "Einstellungen!$B$7:$B$12",
        "NGrenze": "Einstellungen!$C$7:$C$12",
        "NText": "Einstellungen!$D$7:$D$12",
        "PktStd": "Einstellungen!$C$16",
        "SchwelleDB": "Einstellungen!$C$17",
        "GewAn": "Einstellungen!$C$18",
        "GewP": "Einstellungen!$C$19",
        "GewT": "Einstellungen!$C$20",
        "WdhStd": "Einstellungen!$C$21",
        "WarnNote": "Einstellungen!$C$22",
        "Berufe": "Einstellungen!$B$25:$B$34",
    }
    for n, ziel in namen.items():
        wb.defined_names.add(DefinedName(n, attr_text=ziel))


# =============================================================== Stammdaten ==
def baue_stammdaten(ws):
    blattkopf(ws, "L", "Azubi-Stammdaten",
              "Eine Zeile je Azubi, fest mit den Übungsblättern verbunden – "
              "nicht sortieren, nichts löschen, nur den Status ändern")
    breiten(ws, {"A": 6, "B": 18, "C": 16, "D": 32, "E": 10, "F": 16,
                 "G": 15, "H": 10, "I": 10, "J": 26, "K": 30, "L": 3})
    legende_eingabe(ws, "A3:L3")
    hoehen(ws, {3: 24, 4: 10})

    kopf = [("A", "Nr."), ("B", "Name"), ("C", "Vorname"),
            ("D", "Ausbildungsberuf"), ("E", "Jahrgang"),
            ("F", "Ausbildungsstart"), ("G", "Status"), ("H", "Lehrjahr"),
            ("I", "akt. HJ"), ("J", "Anzeigename"), ("K", "Bemerkung")]
    for sp, t in kopf:
        setze(ws, f"{sp}5", t, font=f(10.5, bold=True, color=TXT_WEISS),
              bg=NAVY_500, align=MITTE_UM, border=rahmen(PANEL))
    hoehen(ws, {5: 34})

    for i in range(MAX_AZUBIS):
        r = SD0 + i
        setze(ws, f"A{r}", "=ROW()-5", font=f(10, color=TXT_HELL), bg=FORMEL_BG,
              align=MITTE, border=RAHMEN_ZART)
        for sp in ("B", "C", "D", "F", "G", "K"):
            setze(ws, f"{sp}{r}", None, font=f(11), bg=EINGABE_BG,
                  align=LINKS, border=RAHMEN_EINGABE)
        ws[f"F{r}"].number_format = FMT_DATUM
        ws[f"F{r}"].alignment = MITTE
        ws[f"G{r}"].alignment = MITTE
        setze(ws, f"E{r}", f'=IF($F{r}="","",YEAR($F{r}))',
              font=f(11, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt="0")
        setze(ws, f"H{r}",
              f'=IF(OR($B{r}="",$F{r}=""),"",IF(TODAY()<$F{r},0,'
              f'INT(DATEDIF($F{r},TODAY(),"m")/12)+1))',
              font=f(11, bold=True, color=FORMEL_TXT), bg=FORMEL_BG,
              align=MITTE, border=RAHMEN_ZART, fmt=FMT_HJ)
        setze(ws, f"I{r}",
              f'=IF(OR($B{r}="",$F{r}=""),"",IF(TODAY()<$F{r},0,'
              f'INT(DATEDIF($F{r},TODAY(),"m")/6)+1))',
              font=f(11, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt=FMT_HJ)
        setze(ws, f"J{r}", f'=IF($B{r}="","",$B{r}&", "&$C{r})',
              font=f(11, color=FORMEL_TXT), bg=FORMEL_BG, align=LINKS,
              border=RAHMEN_ZART)
        hoehen(ws, {r: 22})

    dv_beruf = DataValidation(type="list", formula1="Berufe", allow_blank=True,
                              showErrorMessage=True,
                              errorTitle="Beruf nicht bekannt",
                              error="Bitte einen Beruf aus der Liste wählen. "
                                    "Neue Berufe trägt man im Blatt "
                                    "„Einstellungen“ ein.")
    ws.add_data_validation(dv_beruf)
    dv_beruf.add(f"D{SD0}:D{SD1}")

    dv_status = DataValidation(
        type="list", formula1='"aktiv,ausgeschieden,fertig"', allow_blank=True,
        showErrorMessage=True, errorTitle="Status nicht bekannt",
        error="Bitte aktiv, ausgeschieden oder fertig wählen.",
        showInputMessage=True, promptTitle="Status",
        prompt="aktiv = wird in Auswertungen gezeigt. ausgeschieden/fertig = "
               "Daten bleiben erhalten, erscheinen aber nicht mehr in den "
               "Auswahllisten.")
    ws.add_data_validation(dv_status)
    dv_status.add(f"G{SD0}:G{SD1}")

    dv_start = DataValidation(
        type="date", operator="greaterThan", formula1="DATE(1990,1,1)",
        allow_blank=True, showErrorMessage=True,
        errorTitle="Datum nicht lesbar",
        error="Bitte ein Datum im Format TT.MM.JJJJ eingeben, z. B. "
              "01.09.2025.",
        showInputMessage=True, promptTitle="Ausbildungsstart",
        prompt="In der Regel der 01.09. des Einstellungsjahres. Daraus "
               "berechnet die Datei Jahrgang, Lehrjahr und Halbjahr.")
    ws.add_data_validation(dv_start)
    dv_start.add(f"F{SD0}:F{SD1}")

    # Fehlender Ausbildungsstart fällt sofort ins Auge
    ws.conditional_formatting.add(
        f"F{SD0}:F{SD1}",
        FormulaRule(formula=[f'AND($B{SD0}<>"",$F{SD0}="")'],
                    fill=fill("FDE3E1"), font=f(11, bold=True, color="C0392B")))
    # Zwei gleiche Anzeigenamen: die Auswahl trifft sonst immer den ersten
    ws.conditional_formatting.add(
        f"J{SD0}:J{SD1}",
        FormulaRule(formula=[f'AND($J{SD0}<>"",COUNTIF($J${SD0}:$J${SD1},'
                             f'$J{SD0}&"")>1)'],
                    fill=fill("FDE3E1"), font=f(11, bold=True,
                                                color="C0392B"),
                    stopIfTrue=True))
    ws.conditional_formatting.add(
        f"A{SD0}:K{SD1}",
        FormulaRule(formula=[f'$G{SD0}="ausgeschieden"'],
                    font=f(11, italic=True, color=TXT_HELL)))

    verbinde(
    ws, f"A{SD1 + 2}:K{SD1 + 4}",
    "Wichtig: Jede Zeile ist fest mit einer Spalte auf den Übungsblättern "
    "verbunden. Bitte keine Zeilen sortieren, einfügen oder löschen – sonst "
    "wandern vorhandene Noten zum falschen Azubi. Neue Azubis kommen in die "
    "nächste freie Zeile; wer die Ausbildung verlässt, behält seine Zeile und "
    "bekommt nur einen anderen Status. Mehrere Jahrgänge dürfen gemischt in "
    "dieser Liste stehen – auf den Übungsblättern steht über jeder Spalte das "
    "Lehrjahr zum Zeitpunkt der Übung. Mehr als 100 Azubis: siehe Blatt "
    "„Anleitung“, Abschnitt 7.",
    font=f(10, color=TXT), bg=BG, align=LINKS_UM)
    aussenrahmen(ws, f"A{SD1 + 2}:K{SD1 + 4}", LINIE_STARK, "thin")
    hoehen(ws, {SD1 + 2: 20, SD1 + 3: 20, SD1 + 4: 20})
    druck(ws, f"A1:K{SD1 + 3}", quer=True, kopfzeilen="1:5")
    ws.freeze_panes = "B6"


# ============================================================= Übungsliste ===
def baue_uebungsliste(ws):
    blattkopf(ws, "P", "Übungen verwalten",
              "Ein Name in Spalte „Übungsname“ schaltet das zugehörige "
              "Übungsblatt frei – alles Weitere wird dort eingetragen")
    breiten(ws, {"A": 6, "B": 12, "C": 30, "D": 11, "E": 13, "F": 15,
                 "G": 12, "H": 9, "I": 22, "J": 11, "K": 11, "L": 11,
                 "M": 11, "N": 20, "O": 34, "P": 3})
    legende_eingabe(ws, "A3:P3")
    hoehen(ws, {3: 24, 4: 10})

    kopf = [("A", "Nr."), ("B", "Blatt"), ("C", "Übungsname"), ("D", "Art"),
            ("E", "Datum"), ("F", "Status"), ("G", "Bewertet von"),
            ("H", "Gewicht"), ("I", "Wiederholung von"), ("J", "Parameter"),
            ("K", "max. Pkt."), ("L", "bewertet"), ("M", "Ø Prozent"),
            ("N", "Ø Note"), ("O", "Hinweis")]
    for sp, t in kopf:
        setze(ws, f"{sp}5", t, font=f(10.5, bold=True, color=TXT_WEISS),
              bg=NAVY_500, align=MITTE_UM, border=rahmen(PANEL))
    hoehen(ws, {5: 34})

    for i in range(N_UEBUNGEN):
        r = UV0 + i
        blatt = UEB_BLATT.format(i + 1)
        q = f"'{blatt}'"
        setze(ws, f"A{r}", "=ROW()-5", font=f(10, color=TXT_HELL),
              bg=FORMEL_BG, align=MITTE, border=RAHMEN_ZART)
        link_zelle(ws, f"B{r}", f"'{blatt}'!A1", blatt,
                   font=f(10.5, bold=True, color=TEAL, underline="single"),
                   bg=FORMEL_BG, align=MITTE, border=RAHMEN_ZART)
        setze(ws, f"C{r}", None, font=f(11, bold=True), bg=EINGABE_BG,
              align=LINKS, border=RAHMEN_EINGABE)
        setze(ws, f"D{r}", f'=IF($C{r}="","",IF({q}!{Z_ART}="","— fehlt —",'
                           f'{q}!{Z_ART}))',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART)
        setze(ws, f"E{r}", f'=IF(OR($C{r}="",{q}!{Z_DATUM}=""),"",{q}!{Z_DATUM})',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt=FMT_DATUM)
        setze(ws, f"F{r}", f'=IF($C{r}="","",IF({q}!{Z_STATUS}="","— fehlt —",'
                           f'{q}!{Z_STATUS}))',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART)
        setze(ws, f"G{r}", f'=IF($C{r}="","",{q}!{Z_WER})',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART)
        setze(ws, f"H{r}", f'=IF($C{r}="","",IF({q}!{Z_GEWICHT}="",1,{q}!{Z_GEWICHT}))',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt="0.##")
        setze(ws, f"I{r}", f'=IF(OR($C{r}="",{q}!{Z_WDH}=""),"",{q}!{Z_WDH})',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=LINKS,
              border=RAHMEN_ZART)
        setze(ws, f"J{r}", f'=IF($C{r}="","",COUNTIF({q}!{Z_PMAX},">0"))',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt="0")
        setze(ws, f"K{r}", f'=IF($C{r}="","",SUM({q}!{Z_PMAX}))',
              font=f(10.5, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt="0.##")
        setze(ws, f"L{r}", f'=IF($C{r}="","",COUNT({q}!$A$1))',
              font=f(10.5, bold=True, color=FORMEL_TXT), bg=FORMEL_BG,
              align=MITTE, border=RAHMEN_ZART, fmt="0")
        setze(ws, f"M{r}",
              f'=IF(OR($C{r}="",$L{r}=0),"",'
              f'AVERAGE({q}!$A$1))',
              font=f(11, bold=True), bg=PANEL, align=MITTE,
              border=RAHMEN_ZART, fmt=FMT_PROZ)
        setze(ws, f"N{r}", f'=IF($M{r}="","",{note_text(f"$M{r}")})',
              font=f(10.5, bold=True), bg=PANEL, align=MITTE,
              border=RAHMEN_ZART)
        setze(ws, f"O{r}", _hinweis_formel(r, q), font=f(9.5, color=TXT_HELL),
              bg=FORMEL_BG, align=LINKS_UM, border=RAHMEN_ZART)
        hoehen(ws, {r: 22})

    from stil import ampel_regeln
    ampel_regeln(ws, f"M{UV0}:N{UV1}", f"$M{UV0}")
    ws.conditional_formatting.add(
        f"A{UV0}:O{UV1}",
        FormulaRule(formula=[f'$C{UV0}=""'], fill=fill(GESPERRT_BG),
                    font=f(10.5, color="A0A8B4")))

    verbinde(ws, f"A{UV1 + 2}:O{UV1 + 4}",
             "So legen Sie eine neue Übung an: 1. Hier in einer freien Zeile "
             "den Übungsnamen eintragen.  2. Auf den Blattnamen links klicken. "
             " 3. Dort Art, Datum, Status, Kürzel und die Bewertungsparameter "
             "mit Höchstpunktzahl eintragen.  4. Punkte je Azubi eintragen.  "
             "Werden mehr als 20 Übungen gebraucht, das Blatt „Vorlage“ "
             "kopieren (Rechtsklick auf den Blattreiter > Verschieben oder "
             "kopieren > Kopie erstellen).",
             font=f(10, color=TXT), bg=BG, align=LINKS_UM)
    aussenrahmen(ws, f"A{UV1 + 2}:O{UV1 + 4}", LINIE_STARK, "thin")
    hoehen(ws, {UV1 + 2: 22, UV1 + 3: 22, UV1 + 4: 22})
    druck(ws, f"A1:O{UV1 + 4}", quer=True, kopfzeilen="1:5")
    ws.freeze_panes = "C6"


def _hinweis_formel(r, q):
    teile = [
        f'IF({q}!{Z_ART}="","Art fehlt. ","")',
        f'IF({q}!{Z_DATUM}="","Datum fehlt. ","")',
        f'IF({q}!{Z_STATUS}="","Status fehlt. ","")',
        f'IF({q}!{Z_WER}="","Kürzel fehlt. ","")',
        f'IF(COUNTIF({q}!{Z_PMAX},">0")=0,"Parameter fehlen. ","")',
    ]
    inner = "&".join(teile)
    return (f'=IF($C{r}="","frei",IF(LEN({inner})=0,"vollständig",'
            f'"Bitte ergänzen: "&{inner}))')


# =========================================================== Übungsblätter ===
def baue_uebungsblatt(ws, nr):
    """nr = 1..20 für die vorbereiteten Blätter, None für die Vorlage."""
    vorlage = nr is None
    vr = None if vorlage else UV0 + nr - 1
    ws.sheet_properties.tabColor = NAVY_500 if vorlage else TEAL
    ws.sheet_view.showGridLines = False
    breiten(ws, {"A": 4, "B": 22, "C": 24, "D": 9})
    for i in range(MAX_AZUBIS):
        # 15 Einheiten: gerade breit genug für „3 – befriedigend“
        ws.column_dimensions[az_sp(i)].width = 15
    hoehen(ws, {U_TITEL: 22, U_BANNER: 16, U_K1: 18, U_K2: 18, U_NAME: 20,
                U_VORNAME: 15, U_PUNKTE: 15, U_PROZ: 18, U_NOTE: 24,
                U_THEAD: 24})

    _uebung_kopf(ws, nr, vorlage, vr)
    _uebung_azubispalten(ws)
    _uebung_kriterien(ws)
    _uebung_fuss(ws, vorlage)
    _uebung_pruefungen(ws, vorlage)

    # Fixiert bleibt alles bis zum Tabellenkopf: Namen und Ergebnis stehen
    # still, während unten die Kriterien durchgescrollt werden.
    ws.freeze_panes = f"{AZ_ERSTE}{U_P0}"
    # Der Druckbereich deckt bewusst nur den ersten Jahrgangsblock ab:
    # alle 100 Spalten zu drucken ergäbe rund 40 Seiten. Für einen anderen
    # Jahrgang markiert man die Spalten und druckt die Auswahl.
    ws.print_area = f"A1:{az_sp(DRUCK_SPALTEN - 1)}{U_HINW + 1}"
    ws.print_title_rows = f"1:{U_THEAD}"
    ws.print_title_cols = "A:D"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.4


def _uebung_kopf(ws, nr, vorlage, vr):
    titel = "Vorlage für weitere Übungen" if vorlage else f"Übung {nr:02d}"
    verbinde(ws, f"A{U_TITEL}:L{U_TITEL}", titel,
             font=f(16, bold=True, color=TXT_WEISS), bg=NAVY_700,
             align=Alignment(horizontal="left", vertical="center", indent=1))
    bereich(ws, f"A{U_TITEL}:{AZ_LETZTE}{U_TITEL}", bg=NAVY_700)
    verbinde(ws, f"M{U_TITEL}:Q{U_TITEL}")
    link_zelle(ws, f"M{U_TITEL}", "Start!A1", "«  Zurück zur Startseite",
               font=f(10, bold=True, color=TXT_WEISS, underline="single"),
               bg=NAVY_700, align=MITTE)

    if vorlage:
        banner = ('="VORLAGE – dieses Blatt wird in keiner Auswertung '
                  'mitgezählt. Für eine 21. Übung: Blattreiter rechts '
                  'anklicken > Verschieben oder kopieren > Kopie erstellen."')
    else:
        banner = (
            f'=IF({Z_NAME}="",'
            f'"FREI – noch nicht belegt. Zum Freischalten im Blatt '
            f'›Übungen verwalten‹ in Zeile {vr} einen Übungsnamen eintragen.",'
            f'"Belegt · "&{Z_NAME}'
            f'&IF({Z_ART}=""," · Art fehlt"," · "&{Z_ART})'
            f'&IF({Z_DATUM}=""," · Datum fehlt"," · "&'
            f'{datum_text(Z_DATUM)})'
            f'&IF({Z_STATUS}=""," · Status fehlt"," · "&{Z_STATUS})'
            f'&IF(COUNTIF({Z_PMAX},">0")=0," · noch keine Parameter",'
            f'" · "&COUNTIF({Z_PMAX},">0")&" Parameter, max. "'
            f'&SUM({Z_PMAX})&" Punkte"))')
    verbinde(ws, f"A{U_BANNER}:{AZ_LETZTE}{U_BANNER}", banner,
             font=f(9.5, bold=True, color=TXT_WEISS), bg=TEAL,
             align=Alignment(horizontal="left", vertical="center", indent=1))
    if not vorlage:
        ws.conditional_formatting.add(
            f"A{U_BANNER}:{AZ_LETZTE}{U_BANNER}",
            FormulaRule(formula=[f'{Z_NAME}=""'], fill=fill("E8A33D"),
                        font=f(9.5, bold=True, color="3A2A05"),
                        stopIfTrue=True))

    def label(rng, text):
        verbinde(ws, rng, text, font=f(10, bold=True, color=NAVY_700),
                 bg=NAVY_100, align=LINKS, border=RAHMEN_ZART)

    def eingabe(rng, wert=None, fmt=None, align=LINKS):
        verbinde(ws, rng, wert, font=f(10.5, bold=True), bg=EINGABE_BG,
                 align=align, border=RAHMEN_EINGABE, fmt=fmt)

    label(f"A{U_K1}:B{U_K1}", " Übungsname")
    if vorlage:
        eingabe(f"C{U_K1}:E{U_K1}")
    else:
        verbinde(ws, f"C{U_K1}:E{U_K1}",
                 f"=IF('Übungen'!$C{vr}=\"\",\"\",'Übungen'!$C{vr})",
                 font=f(11, bold=True, color=FORMEL_TXT), bg=FORMEL_BG,
                 align=LINKS, border=RAHMEN_ZART)
    label(f"F{U_K1}:G{U_K1}", " Art")
    eingabe(f"H{U_K1}:I{U_K1}", align=MITTE)
    label(f"J{U_K1}:K{U_K1}", " Datum")
    eingabe(f"L{U_K1}:M{U_K1}", fmt=FMT_DATUM, align=MITTE)
    verbinde(ws, f"N{U_K1}:{AZ_LETZTE}{U_K1}",
             "Bewertungsparameter untereinander, Azubis nebeneinander – zum "
             "Eintragen nach unten scrollen.", font=f(9, color=TXT_HELL), align=LINKS)

    label(f"A{U_K2}:B{U_K2}", " Status")
    eingabe(f"C{U_K2}:E{U_K2}")
    label(f"F{U_K2}:G{U_K2}", " Bewertet von")
    eingabe(f"H{U_K2}:I{U_K2}", align=MITTE)
    label(f"J{U_K2}:K{U_K2}", " Gewicht")
    eingabe(f"L{U_K2}:M{U_K2}", 1, fmt="0.##", align=MITTE)
    label(f"N{U_K2}:P{U_K2}", " Wiederholung von")
    eingabe(f"Q{U_K2}:S{U_K2}")
    verbinde(ws, f"T{U_K2}:{AZ_LETZTE}{U_K2}",
             "Nur ausfüllen, wenn dies der Nachhol- oder Wiederholtermin "
             "einer anderen Übung ist.", font=f(9, color=TXT_HELL),
             align=LINKS)


def _uebung_azubispalten(ws):
    """Kopfzeilen je Azubi: Name, Vorname und das laufende Ergebnis."""
    zeilen = [(U_NAME, " Name"), (U_VORNAME, " Vorname · Lehrjahr"),
              (U_PUNKTE, " Punkte"), (U_PROZ, " Prozent"),
              (U_NOTE, " IHK-Note")]
    for r, text in zeilen:
        verbinde(ws, f"A{r}:D{r}", text,
                 font=f(10, bold=(r in (U_NAME, U_PROZ, U_NOTE)),
                        color=NAVY_700), bg=NAVY_100, align=LINKS,
                 border=RAHMEN_ZART)

    for i in range(MAX_AZUBIS):
        c = az_sp(i)
        sd = SD0 + i
        leer = f'IF(\'Stammdaten\'!$B{sd}="","",'
        setze(ws, f"{c}{U_NAME}", f"={leer}'Stammdaten'!$B{sd})",
              font=f(11, bold=True, color=FORMEL_TXT), bg=FORMEL_BG,
              align=MITTE_UM, border=RAHMEN_ZART)
        # Vorname und das Lehrjahr, in dem der Azubi zum Zeitpunkt DIESER
        # Übung stand – nicht das heutige. So lassen sich mehrere Jahrgänge
        # in einer Liste führen, ohne die Spalten zu verwechseln.
        setze(ws, f"{c}{U_VORNAME}",
              f'={leer}\'Stammdaten\'!$C{sd}&IF({c}${U_HJ}>0," · LJ "&'
              f'(INT(({c}${U_HJ}-1)/2)+1),""))',
              font=f(10, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART)
        setze(ws, f"{c}{U_PUNKTE}",
              f'=IF({c}${U_SUMME}="","",{c}${U_SUMME}&" / "&{c}${U_MAX})',
              font=f(10, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART)
        setze(ws, f"{c}{U_PROZ}",
              f'=IF(OR({c}${U_SUMME}="",{c}${U_MAX}=""),"",'
              f'IF({c}${U_MAX}=0,"",{c}${U_SUMME}/{c}${U_MAX}))',
              font=f(12, bold=True), bg=PANEL, align=MITTE,
              border=RAHMEN_ZART, fmt=FMT_PROZ)
        setze(ws, f"{c}{U_NOTE}",
              f'=IF({c}${U_NAME}="","",'
              f'IF(COUNT({c}${U_P0}:{c}${U_P1})>0,'
              f'{note_text(f"{c}${U_PROZ}")},'
              f'IF(COUNTIF({c}${U_P0}:{c}${U_P1},"n.b.")>0,"n.b."&'
              f'IF({c}${U_GRUND}=""," (Grund fehlt)",'
              f'" ("&{c}${U_GRUND}&")"),"")))',
              font=f(10.5, bold=True), bg=PANEL, align=MITTE_UM,
              border=RAHMEN_ZART)


def _uebung_kriterien(ws):
    """Tabellenkopf und die Kriterienzeilen mit den Punktefeldern."""
    for sp, t in (("A", "Nr."), ("B", "Bewertungsparameter"),
                  ("C", "Bewertungshorizont"), ("D", "max. Pkt.")):
        setze(ws, f"{sp}{U_THEAD}", t,
              font=f(10, bold=True, color=TXT_WEISS), bg=NAVY_700,
              align=MITTE_UM, border=rahmen(PANEL))
    verbinde(ws, f"{AZ_ERSTE}{U_THEAD}:{AZ_LETZTE}{U_THEAD}",
             "Punkte je Azubi eintragen – „n.b.“ für nicht bewertet",
             font=f(10, bold=True, color=TXT_WEISS), bg=NAVY_700,
             align=MITTE, border=rahmen(PANEL))

    for k in range(MAX_PARAM):
        r = U_P0 + k
        setze(ws, f"A{r}", f'=IF($B{r}="","",ROW()-{U_P0 - 1})',
              font=f(9.5, color=TXT_HELL), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART)
        setze(ws, f"B{r}", None, font=f(10, bold=True), bg=EINGABE_BG,
              align=LINKS_UM, border=RAHMEN_EINGABE)
        setze(ws, f"C{r}", None, font=f(8.5, color=TXT_HELL), bg=EINGABE_BG,
              align=LINKS_UM, border=RAHMEN_EINGABE)
        setze(ws, f"D{r}", f'=IF($B{r}="",0,PktStd)',
              font=f(10, bold=True, color=FORMEL_TXT), bg=FORMEL_BG,
              align=MITTE, border=RAHMEN_ZART, fmt='0.##;-0.##;""')
        for i in range(MAX_AZUBIS):
            setze(ws, f"{az_sp(i)}{r}", None, font=f(11), bg=EINGABE_BG,
                  align=MITTE, border=RAHMEN_EINGABE, fmt=FMT_PUNKT)
        hoehen(ws, {r: 24})

    # Werte, die unter dem Kriterienblock stehen: Rechenergebnisse und die
    # Felder, die nur im Einzelfall gebraucht werden.
    unten = [
        (U_SUMME, " Summe", True), (U_MAX, " von max.", True),
        (U_HJ, " Halbjahr", True), (U_GRUND, " Grund (n.b.)", False),
        (U_BEM, " Bemerkung", False),
        (U_WERT, " Wertung bei Wiederholung", False),
        (U_STAT, " Status des Azubis", True),
    ]
    for r, text, berechnet in unten:
        verbinde(ws, f"A{r}:D{r}", text,
                 font=f(9.5 if r == U_STAT else 10, bold=(r == U_SUMME),
                        color=TXT_HELL if r == U_STAT else NAVY_700),
                 bg=BG if r == U_STAT else NAVY_100, align=LINKS,
                 border=RAHMEN_ZART)
        hoehen(ws, {r: 34 if r == U_BEM else 20})

    for i in range(MAX_AZUBIS):
        c = az_sp(i)
        sd = SD0 + i
        setze(ws, f"{c}{U_SUMME}",
              f'=IF(OR({c}${U_NAME}="",COUNT({c}${U_P0}:{c}${U_P1})=0),"",'
              f'SUM({c}${U_P0}:{c}${U_P1}))',
              font=f(11, bold=True, color=FORMEL_TXT), bg=FORMEL_BG,
              align=MITTE, border=RAHMEN_ZART, fmt=FMT_PUNKT)
        setze(ws, f"{c}{U_MAX}",
              f'=IF({c}${U_SUMME}="","",SUMPRODUCT('
              f'--ISNUMBER({c}${U_P0}:{c}${U_P1}),{Z_PMAX}))',
              font=f(10, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt=FMT_PUNKT)
        setze(ws, f"{c}{U_HJ}",
              f'=IF({c}${U_NAME}="","",'
              f'IF(OR({Z_DATUM}="",\'Stammdaten\'!$F{sd}=""),0,'
              f'IF({Z_DATUM}<\'Stammdaten\'!$F{sd},0,'
              f'INT(DATEDIF(\'Stammdaten\'!$F{sd},{Z_DATUM},"m")/6)+1)))',
              font=f(10, color=FORMEL_TXT), bg=FORMEL_BG, align=MITTE,
              border=RAHMEN_ZART, fmt=FMT_HJ)
        setze(ws, f"{c}{U_GRUND}", None, font=f(10), bg=EINGABE_BG,
              align=MITTE, border=RAHMEN_EINGABE)
        setze(ws, f"{c}{U_BEM}", None, font=f(9), bg=EINGABE_BG,
              align=LINKS_UM, border=RAHMEN_EINGABE)
        setze(ws, f"{c}{U_WERT}", None, font=f(10), bg=EINGABE_BG,
              align=MITTE, border=RAHMEN_EINGABE)
        # Nur für die bedingte Formatierung: quer über Spalten verschiebt
        # sich ein Bezug auf die Stammdaten nicht mit der Zeile.
        setze(ws, f"{c}{U_STAT}",
              f'=IF(\'Stammdaten\'!$B{sd}="","",\'Stammdaten\'!$G{sd})',
              font=f(8.5, color=TXT_HELL), bg=BG, align=MITTE)


def _uebung_fuss(ws, vorlage):
    legende_eingabe(ws, f"A{U_LEG}:{AZ_LETZTE}{U_LEG}")
    hoehen(ws, {U_LEG: 22, U_FARB: 22})
    verbinde(ws, f"A{U_FARB}:B{U_FARB}", " Farblegende",
             font=f(10, bold=True, color=NAVY_700), align=LINKS)
    notenlegende_bereiche(
        ws, [f"C{U_FARB}:D{U_FARB}", f"E{U_FARB}:F{U_FARB}",
             f"G{U_FARB}:H{U_FARB}", f"I{U_FARB}:J{U_FARB}",
             f"K{U_FARB}:L{U_FARB}", f"M{U_FARB}:N{U_FARB}"])
    verbinde(ws, f"O{U_FARB}:{AZ_LETZTE}{U_FARB}",
             f"Farbe und Zahl sagen dasselbe – der Ausdruck in Schwarz-Weiß "
             f"bleibt lesbar. Gedruckt werden die ersten {DRUCK_SPALTEN} "
             f"Azubi-Spalten; für einen anderen Jahrgang die Spalten "
             f"markieren und „Auswahl drucken“ wählen.",
             font=f(9, color=TXT_HELL), align=LINKS_UM)
    verbinde(
        ws, f"A{U_HINW}:{AZ_LETZTE}{U_HINW + 1}",
        "Nicht bewertet: Einzelne Punktefelder mit „n.b.“ fallen aus Summe "
        "und Höchstpunktzahl heraus. Steht in der ganzen Spalte eines Azubis "
        "nur „n.b.“, zählt die Übung für ihn nirgends mit – weder als "
        "0 Punkte noch im Durchschnitt, in der Halbjahresnote oder in der "
        "Warnübersicht. Bitte zusätzlich unten in der Zeile „Grund (n.b.)“ "
        "den Grund wählen.",
        font=f(10, color=TXT), bg=BG, align=LINKS_UM)
    aussenrahmen(ws, f"A{U_HINW}:{AZ_LETZTE}{U_HINW + 1}", LINIE_STARK,
                 "thin")
    hoehen(ws, {U_HINW: 20, U_HINW + 1: 20})


def _uebung_pruefungen(ws, vorlage):
    """Eingabeprüfung und bedingte Formatierung eines Übungsblattes."""
    E, L = AZ_ERSTE, AZ_LETZTE

    def dv(**kw):
        d = DataValidation(**kw)
        ws.add_data_validation(d)
        return d

    dv(type="list", formula1='"Praxis,Theorie"', allow_blank=True,
       showErrorMessage=True, errorTitle="Art nicht bekannt",
       error="Bitte Praxis oder Theorie wählen.", showInputMessage=True,
       promptTitle="Art der Übung",
       prompt="Praxis = praktische Übung, Theorie = schriftlicher Test."
       ).add(Z_ART.replace("$", ""))
    dv(type="list", formula1='"laufend,abgeschlossen"', allow_blank=True,
       showErrorMessage=True, errorTitle="Status nicht bekannt",
       error="Bitte laufend oder abgeschlossen wählen."
       ).add(Z_STATUS.replace("$", ""))
    dv(type="date", operator="greaterThan", formula1="DATE(1990,1,1)",
       allow_blank=True, showErrorMessage=True,
       errorTitle="Datum nicht lesbar",
       error="Bitte ein Datum im Format TT.MM.JJJJ eingeben, z. B. "
             "12.03.2026.").add(Z_DATUM.replace("$", ""))
    dv(type="decimal", operator="between", formula1="0.1", formula2="10",
       allow_blank=True, showErrorMessage=True,
       errorTitle="Gewicht außerhalb des Bereichs",
       error="Bitte eine Zahl zwischen 0,1 und 10 eingeben. 1 = normal."
       ).add(Z_GEWICHT.replace("$", ""))
    if not vorlage:
        dv(type="list", formula1=f"'Übungen'!$C${UV0}:$C${UV1}",
           allow_blank=True, showErrorMessage=True,
           errorTitle="Übung nicht gefunden",
           error="Bitte eine Übung aus der Liste wählen oder das Feld leer "
                 "lassen.").add(Z_WDH.replace("$", ""))
    dv(type="decimal", operator="between", formula1="0", formula2="1000",
       allow_blank=True, showErrorMessage=True,
       errorTitle="Höchstpunktzahl nicht gültig",
       error="Bitte eine Zahl zwischen 0 und 1000 eingeben."
       ).add(f"D{U_P0}:D{U_P1}")
    dv(type="custom",
       formula1=f'OR(EXACT({E}{U_P0},"n.b."),AND(ISNUMBER({E}{U_P0}),'
                f'{E}{U_P0}>=0,{E}{U_P0}<=$D{U_P0}))',
       allow_blank=True, showErrorMessage=True,
       errorTitle="Eingabe nicht möglich",
       error="Erlaubt sind: eine Zahl von 0 bis zur Höchstpunktzahl in "
             "Spalte D – oder n.b. für „nicht bewertet“.\n"
             "Steht in Spalte D noch keine Punktzahl, bitte zuerst links in "
             "dieser Zeile den Parameter eintragen.",
       showInputMessage=True, promptTitle="Punkte eintragen",
       prompt="Zahl eingeben – oder n.b., wenn der Azubi hier nicht bewertet "
              "wurde.").add(f"{E}{U_P0}:{L}{U_P1}")
    dv(type="list", formula1='"krank,Urlaub,Berufsschule,sonstiges"',
       allow_blank=True, showErrorMessage=True,
       errorTitle="Grund nicht bekannt",
       error="Bitte krank, Urlaub, Berufsschule oder sonstiges wählen.",
       showInputMessage=True, promptTitle="Grund für n.b.",
       prompt="Nur ausfüllen, wenn der Azubi bei dieser Übung nicht bewertet "
              "wurde.").add(f"{E}{U_GRUND}:{L}{U_GRUND}")
    dv(type="list", formula1='"letzte,beste,keine"', allow_blank=True,
       showErrorMessage=True, errorTitle="Wertung nicht bekannt",
       error="Bitte letzte, beste oder keine wählen.", showInputMessage=True,
       promptTitle="Wertung bei Wiederholung",
       prompt="Nur auf einem Wiederholungsblatt nötig. Leer = Standard aus "
              "den Einstellungen.").add(f"{E}{U_WERT}:{L}{U_WERT}")

    # Reihenfolge der Regeln ist Absicht: „frei“ und „keine Spalte“ decken
    # alles darunter ab und halten die übrigen Regeln an.
    if not vorlage:
        ws.conditional_formatting.add(
            f"A{U_NAME}:{L}{U_WERT}",
            FormulaRule(formula=[f'{Z_NAME}=""'], fill=fill(GESPERRT_BG),
                        font=f(10, color="A6ADB8"), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"{E}{U_NAME}:{L}{U_WERT}",
        FormulaRule(formula=[f'{E}${U_NAME}=""'], fill=fill(GESPERRT_BG),
                    font=f(10, color="C2C8D0"), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"A{U_P0}:{L}{U_P1}",
        FormulaRule(formula=[f'$B{U_P0}=""'], fill=fill(GESPERRT_BG),
                    font=f(10, color="C2C8D0"), stopIfTrue=True))
    ampel_regeln(ws, f"{E}{U_PROZ}:{L}{U_NOTE}", f"{E}${U_PROZ}")
    ws.conditional_formatting.add(
        f"{E}{U_P0}:{L}{U_P1}",
        FormulaRule(formula=[f'EXACT({E}{U_P0},"n.b.")'], fill=fill("E3E7ED"),
                    font=f(10, italic=True, color=TXT_HELL),
                    stopIfTrue=True))
    ws.conditional_formatting.add(
        f"{E}{U_P0}:{L}{U_P1}",
        FormulaRule(formula=[f'AND(ISNUMBER({E}{U_P0}),{E}{U_P0}>$D{U_P0})'],
                    fill=fill("FDE3E1"),
                    font=f(11, bold=True, color="C0392B")))
    # Spalte angefangen, aber noch nicht fertig
    ws.conditional_formatting.add(
        f"{E}{U_NAME}:{L}{U_VORNAME}",
        FormulaRule(
            formula=[f'AND({E}${U_NAME}<>"",'
                     f'COUNT({E}${U_P0}:{E}${U_P1})'
                     f'+COUNTIF({E}${U_P0}:{E}${U_P1},"n.b.")>0,'
                     f'COUNT({E}${U_P0}:{E}${U_P1})'
                     f'+COUNTIF({E}${U_P0}:{E}${U_P1},"n.b.")'
                     f'<COUNTIF({Z_PMAX},">0"))'],
            fill=fill("FFF1D6")))
    # Nicht mehr aktive Azubis zurückhaltend darstellen – der Status steht
    # dafür eigens in der Zeile U_STAT, weil ein Bezug auf die Stammdaten
    # sich quer über Spalten nicht mitverschieben würde.
    ws.conditional_formatting.add(
        f"{E}{U_NAME}:{L}{U_VORNAME}",
        FormulaRule(formula=[f'AND({E}${U_NAME}<>"",'
                             f'{E}${U_STAT}<>"aktiv")'],
                    font=f(10, italic=True, color="98A1AE")))
